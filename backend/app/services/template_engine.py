"""Generic, data-driven DOCX pair validation and one-pass rendering."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date
from pathlib import PurePosixPath

from docx import Document
from docxtpl import DocxTemplate
from lxml import etree
from jinja2 import StrictUndefined, nodes
from jinja2.sandbox import SandboxedEnvironment
from jsonschema import Draft202012Validator

from ..errors import AppError
from ..settings import CONFIG_MAX, DOCX_MAX, VALUES_MAX

ROOT_TOKEN = re.compile(r"\{\{\s*([A-Z][A-Z0-9_]{0,47})\s*\}\}")
ROW_TOKEN = re.compile(r"\{\{\s*row\.([A-Z][A-Z0-9_]{0,47})\s*\}\}")
P_START = re.compile(r"\{%\s*p\s+for\s+row\s+in\s+([A-Z][A-Z0-9_]{0,47})\s*%\}")
P_END = re.compile(r"\{%\s*p\s+endfor\s*%\}")
TR_START = re.compile(r"\{%\s*tr\s+for\s+row\s+in\s+([A-Z][A-Z0-9_]{0,47})\s*%\}")
TR_END = re.compile(r"\{%\s*tr\s+endfor\s*%\}")
ANY_JINJA = re.compile(r"(\{\{.*?\}\}|\{%.*?%\}|\{#.*?#\})", re.S)
ID = re.compile(r"^[a-z][a-z0-9_]{0,47}$")
PH = re.compile(r"^[A-Z][A-Z0-9_]{0,47}$")
SAFE_FONT = re.compile(r"^[\w -]{1,64}$", re.UNICODE)
FORBIDDEN_IDS = {"__proto__", "prototype", "constructor"}
MAX_EXPANDED = 20_971_520
MAX_ENTRY = 5_242_880
MAX_ENTRIES = 500
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _duplicates(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("duplicate")
        result[key] = value
    return result


def parse_json(data, cap, code="TEMPLATE_JSON_INVALID"):
    if not data or len(data) > cap:
        raise AppError("INPUT_TOO_LARGE" if len(data) > cap else code, 413 if len(data) > cap else 422)
    try:
        text = data.decode("utf-8", "strict")
        value = json.loads(text, object_pairs_hook=_duplicates, parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError):
        raise AppError(code, 422)
    if _depth(value) > 16:
        raise AppError(code, 422)
    return value


def _depth(value):
    if isinstance(value, dict): return 1 + max((_depth(v) for v in value.values()), default=0)
    if isinstance(value, list): return 1 + max((_depth(v) for v in value), default=0)
    if isinstance(value, float) and not math.isfinite(value): return 99
    return 0


def fingerprint(docx, config):
    h = hashlib.sha256()
    h.update(b"CLS-TEMPLATE-PAIR-V1\0")
    h.update(len(docx).to_bytes(8, "big")); h.update(docx)
    h.update(len(config).to_bytes(8, "big")); h.update(config)
    return h.hexdigest()


def _schema():
    from pathlib import Path
    return json.loads((Path(__file__).parents[1] / "schemas/template-config.schema.json").read_text("utf-8"))


def _issue(location, code):
    return {"location": location, "code": code}


def validate_config(config):
    errors = sorted(Draft202012Validator(_schema()).iter_errors(config), key=lambda e: list(e.path))
    if errors:
        if config.get("schema_version") != "1.0": raise AppError("TEMPLATE_SCHEMA_UNSUPPORTED", 422)
        raise AppError("TEMPLATE_CONFIG_INVALID", 422, [_issue(".".join(map(str, e.path)) or "configuration", "INVALID") for e in errors[:20]])
    ids, placeholders = set(), set()
    for field in config["fields"]:
        _unique(field["id"], ids); _unique(field["placeholder"], placeholders); _field_definition(field)
    for section in config["sections"]:
        _unique(section["id"], ids); _unique(section["placeholder"], placeholders)
        if section["min_rows"] > section["max_rows"] or len(section["default_rows"]) > section["max_rows"]:
            raise AppError("TEMPLATE_CONFIG_INVALID", 422)
        rowids, rowph = set(), set()
        for field in section["row_fields"]:
            _unique(field["id"], rowids); _unique(field["placeholder"], rowph); _field_definition(field)
        if section["numbering"]["placeholder"] in rowph:
            raise AppError("TEMPLATE_CONFIG_INVALID", 422)
        for row in section["default_rows"]:
            if set(row) != rowids: raise AppError("TEMPLATE_CONFIG_INVALID", 422)
            for field in section["row_fields"]: _field_value(field, row[field["id"]], defaults=True)
    if not SAFE_FONT.fullmatch(config["preview"]["font_family"]): raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    for block in config["preview"]["blocks"]:
        if block["type"] == "table" and abs(sum(c["width_percent"] for c in block["columns"]) - 100) > .01:
            raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    for pattern, suffix in ((config["output"]["docx_filename_pattern"], ".docx"),
                            (config["output"]["transcript_filename_pattern"], ".txt")):
        if not pattern.lower().endswith(suffix) or any(c in pattern for c in "/\\\r\n\0") or ROW_TOKEN.search(pattern):
            raise AppError("OUTPUT_NAME_INVALID", 422)
        _tokens_allowed(pattern, placeholders, set())
    return config


def _unique(value, seen):
    if value in seen or value in FORBIDDEN_IDS or value.startswith("_"):
        raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    seen.add(value)


def _field_definition(field):
    if field["id"] in FORBIDDEN_IDS or field["id"].startswith("_"): raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    if field["type"] in ("text", "textarea") and field["validation"]["min_length"] > field["validation"]["max_length"]:
        raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    if field["type"] == "integer" and field["validation"]["minimum"] > field["validation"]["maximum"]:
        raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    if field["type"] == "select":
        vals = [x["value"] for x in field["options"]]
        if len(vals) != len(set(vals)): raise AppError("TEMPLATE_CONFIG_INVALID", 422)
    _field_value(field, field["default"], defaults=True)


def _xml_text_valid(value):
    return all(c in "\t\n" or ord(c) >= 32 and not 0xD800 <= ord(c) <= 0xDFFF and ord(c) not in (0xFFFE, 0xFFFF) for c in value)


def _field_value(field, value, defaults=False):
    kind = field["type"]
    if kind in ("text", "textarea"):
        if not isinstance(value, str) or not _xml_text_valid(value): raise AppError("FIELD_INVALID", 422)
        value = value.replace("\r\n", "\n").replace("\r", "\n")
        length = len(value)
        if length > field["validation"]["max_length"] or (value and length < field["validation"]["min_length"]): raise AppError("FIELD_INVALID", 422)
        if not defaults and field["required"] and not value.strip(): raise AppError("FIELD_INVALID", 422)
    elif kind == "integer":
        if value is not None and (isinstance(value, bool) or not isinstance(value, int) or not field["validation"]["minimum"] <= value <= field["validation"]["maximum"]): raise AppError("FIELD_INVALID", 422)
        if not defaults and field["required"] and value is None: raise AppError("FIELD_INVALID", 422)
    elif kind == "date":
        if not isinstance(value, str): raise AppError("FIELD_INVALID", 422)
        if value:
            try: date.fromisoformat(value)
            except ValueError: raise AppError("FIELD_INVALID", 422)
        elif not defaults and field["required"]: raise AppError("FIELD_INVALID", 422)
    elif kind == "select":
        if not isinstance(value, str) or value and value not in {o["value"] for o in field["options"]}: raise AppError("FIELD_INVALID", 422)
        if not defaults and field["required"] and not value: raise AppError("FIELD_INVALID", 422)
    elif kind == "boolean":
        if value is not None and type(value) is not bool: raise AppError("FIELD_INVALID", 422)
        if not defaults and field["required"] and value is None: raise AppError("FIELD_INVALID", 422)
    return value


def _inspect_zip(docx):
    if not docx or len(docx) > DOCX_MAX or not zipfile.is_zipfile(io.BytesIO(docx)):
        raise AppError("INPUT_TOO_LARGE" if len(docx) > DOCX_MAX else "TEMPLATE_ARCHIVE_INVALID", 413 if len(docx) > DOCX_MAX else 415)
    try:
        with zipfile.ZipFile(io.BytesIO(docx)) as z:
            infos = z.infolist()
            if len(infos) > MAX_ENTRIES or len({i.filename for i in infos}) != len(infos): raise AppError("TEMPLATE_TOO_COMPLEX", 422)
            total = 0
            for info in infos:
                name = info.filename
                path = PurePosixPath(name)
                if (path.is_absolute() or "\\" in name or "\0" in name or ".." in path.parts or info.flag_bits & 1
                    or info.file_size > MAX_ENTRY or info.compress_size and info.file_size / info.compress_size > 100):
                    raise AppError("TEMPLATE_UNSAFE_CONTENT", 422)
                total += info.file_size
                if total > MAX_EXPANDED: raise AppError("TEMPLATE_TOO_COMPLEX", 422)
                lower = name.lower()
                if any(x in lower for x in ("vbaproject", "activex", "embeddings/", "customxml/", "comments", "footnotes", "endnotes")) or lower.startswith("word/media/"):
                    raise AppError("TEMPLATE_UNSAFE_CONTENT", 422)
                content = z.read(info)
                if lower.endswith((".xml", ".rels")):
                    if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper(): raise AppError("TEMPLATE_UNSAFE_CONTENT", 422)
                    root = fromstring(content)
                    if lower.endswith(".rels"):
                        for rel in root:
                            if rel.get("TargetMode") == "External": raise AppError("TEMPLATE_UNSAFE_CONTENT", 422)
                    if any(el.tag in {W+"instrText", W+"fldSimple", W+"altChunk", W+"sdt", W+"ins", W+"del", W+"moveFrom", W+"moveTo"} for el in root.iter()):
                        raise AppError("TEMPLATE_UNSAFE_CONTENT", 422)
                    if any(el.tag == W+"vanish" for el in root.iter()): raise AppError("TEMPLATE_UNSAFE_CONTENT", 422)
            required = {"[Content_Types].xml", "word/document.xml"}
            if not required.issubset(z.namelist()): raise AppError("TEMPLATE_ARCHIVE_INVALID", 415)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        if isinstance(exc, AppError): raise
        raise AppError("TEMPLATE_ARCHIVE_INVALID", 415)


def _all_text_parts(docx):
    result = {}
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        for name in z.namelist():
            if name == "word/document.xml" or re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                root = fromstring(z.read(name)); result[name] = root
    return result


def _tokens_allowed(text, roots, rows):
    found = ANY_JINJA.findall(text)
    consumed = []
    for token in found:
        match = ROOT_TOKEN.fullmatch(token)
        if match and match.group(1) in roots: consumed.append(token); continue
        match = ROW_TOKEN.fullmatch(token)
        if match and match.group(1) in rows: consumed.append(token); continue
        if P_START.fullmatch(token) or P_END.fullmatch(token) or TR_START.fullmatch(token) or TR_END.fullmatch(token): consumed.append(token); continue
        raise AppError("TEMPLATE_BINDING_MISMATCH", 422)
    residue = ANY_JINJA.sub("", text)
    if "{{" in residue or "{%" in residue or "{#" in residue or "}}" in residue or "%}" in residue:
        raise AppError("TEMPLATE_PLACEHOLDER_SPLIT", 422)
    return found


def validate_pair(docx: bytes, config_bytes: bytes):
    _inspect_zip(docx)
    config = validate_config(parse_json(config_bytes, CONFIG_MAX))
    roots = {f["placeholder"] for f in config["fields"]}
    sections = {s["placeholder"]: s for s in config["sections"]}
    seen_roots, seen_sections = set(), set()
    for name, root in _all_text_parts(docx).items():
        for paragraph in root.iter(W+"p"):
            runs = ["".join(t.text or "" for t in run.iter(W+"t")) for run in paragraph.iter(W+"r")]
            combined = "".join(runs)
            if any(mark in combined for mark in ("{{", "{%", "{#")):
                for token in ANY_JINJA.findall(combined):
                    if not any(token in run for run in runs): raise AppError("TEMPLATE_PLACEHOLDER_SPLIT", 422, [_issue(name, "SPLIT_TAG")])
                row_allowed = set().union(*({f["placeholder"] for f in s["row_fields"]} | {s["numbering"]["placeholder"]} for s in config["sections"]), set())
                _tokens_allowed(combined, roots, row_allowed)
                seen_roots |= set(ROOT_TOKEN.findall(combined))
                for pattern in (P_START, TR_START):
                    m = pattern.fullmatch(combined.strip())
                    if m: seen_sections.add(m.group(1))
    preview_roots, preview_sections = set(), set()
    for block in config["preview"]["blocks"]:
        if block["type"] in ("paragraph", "heading"):
            _tokens_allowed(block["text"], roots, set()); preview_roots |= set(ROOT_TOKEN.findall(block["text"]))
        elif block["type"] in ("repeat", "table"):
            section = sections.get(block["section"])
            if not section: raise AppError("TEMPLATE_BINDING_MISMATCH", 422)
            preview_sections.add(block["section"])
            allowed = {f["placeholder"] for f in section["row_fields"]}
            if section["numbering"]["style"] != "none": allowed.add(section["numbering"]["placeholder"])
            texts = [x["text"] for x in block.get("blocks", block.get("columns", []))]
            for text in texts: _tokens_allowed(text, set(), allowed)
    if roots - seen_roots or roots - preview_roots or set(sections) != seen_sections or set(sections) != preview_sections:
        raise AppError("TEMPLATE_BINDING_MISMATCH", 422)
    if _docx_content_tree(docx) != _preview_content_tree(config):
        raise AppError("TEMPLATE_PREVIEW_MISMATCH", 422)
    _validate_geometry(docx, config)
    warnings = [{"code": "PREVIEW_LAYOUT_APPROXIMATE", "message": "The preview may differ from the Word document."}]
    if config["status"] == "example": warnings.insert(0, {"code": "EXAMPLE_TEMPLATE", "message": "This example is not approved for filing."})
    return {"fingerprint": fingerprint(docx, config_bytes), "schema_version": "1.0", "configuration": config, "warnings": warnings}


def _canonical(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = ROOT_TOKEN.sub(lambda m: "{{ "+m.group(1)+" }}", text)
    text = ROW_TOKEN.sub(lambda m: "{{ row."+m.group(1)+" }}", text)
    text = P_START.sub(lambda m: "{%p for row in "+m.group(1)+" %}", text)
    text = P_END.sub("{%p endfor %}", text)
    text = TR_START.sub(lambda m: "{%tr for row in "+m.group(1)+" %}", text)
    return TR_END.sub("{%tr endfor %}", text)


def _paragraph_text(element):
    return _canonical("".join(node.text or "" for node in element.iter(W+"t")))


def _part_nodes(root):
    container = root.find(W+"body") if root.tag == W+"document" else root
    result=[]
    for child in container:
        if child.tag == W+"p":
            text=_paragraph_text(child)
            if text: result.append(("p",text))
        elif child.tag == W+"tbl":
            rows=[]
            for tr in child.findall(W+"tr"):
                rows.append(tuple(_paragraph_text(tc) for tc in tr.findall(W+"tc")))
            if len(rows)!=4 or not rows[1] or not TR_START.fullmatch(rows[1][0]) or not rows[3] or not TR_END.fullmatch(rows[3][0]) or any(rows[1][1:]) or any(rows[3][1:]):
                raise AppError("TEMPLATE_LOOP_INVALID",422)
            result.append(("table",TR_START.fullmatch(rows[1][0]).group(1),rows[0],rows[2]))
    return result


def _docx_content_tree(docx):
    parts=_all_text_parts(docx);result=[]
    for name in sorted((n for n in parts if "/header" in n)):
        result.extend(_part_nodes(parts[name]))
    result.extend(_part_nodes(parts["word/document.xml"]))
    for name in sorted((n for n in parts if "/footer" in n)):
        result.extend(_part_nodes(parts[name]))
    return result


def _preview_content_tree(config):
    result=[]
    for block in config["preview"]["blocks"]:
        if block["type"] in ("paragraph","heading"):
            if block["text"]: result.append(("p",_canonical(block["text"])))
        elif block["type"]=="repeat":
            result.append(("p",f"{{%p for row in {block['section']} %}}"))
            result.extend(("p",_canonical(item["text"])) for item in block["blocks"] if item["text"])
            result.append(("p","{%p endfor %}"))
        else:
            result.append(("table",block["section"],tuple(c["label"] for c in block["columns"]),tuple(_canonical(c["text"]) for c in block["columns"])))
    return result


def _validate_geometry(docx, config):
    document = Document(io.BytesIO(docx))
    if len(document.sections) != 1: raise AppError("TEMPLATE_PREVIEW_MISMATCH", 422)
    section = document.sections[0]; page = config["preview"]["page"]
    expected = (210, 297) if page["size"] == "A4" else (215.9, 279.4)
    if page["orientation"] == "landscape": expected = expected[::-1]
    actual = (section.page_width.mm, section.page_height.mm)
    margins = page["margins_mm"]
    checks = zip(actual, expected)
    checks = list(checks) + [(section.top_margin.mm, margins["top"]), (section.right_margin.mm, margins["right"]),
        (section.bottom_margin.mm, margins["bottom"]), (section.left_margin.mm, margins["left"])]
    if any(abs(a-b) > 0.02 for a,b in checks): raise AppError("TEMPLATE_PREVIEW_MISMATCH", 422)


def validate_values(config, envelope):
    if not isinstance(envelope, dict) or set(envelope) != {"fields", "sections"} or not isinstance(envelope["fields"], dict) or not isinstance(envelope["sections"], dict):
        raise AppError("FIELD_INVALID", 422)
    if set(envelope["fields"]) != {f["id"] for f in config["fields"]} or set(envelope["sections"]) != {s["id"] for s in config["sections"]}:
        raise AppError("FIELD_INVALID", 422)
    for f in config["fields"]: _field_value(f, envelope["fields"][f["id"]])
    total = 0
    for section in config["sections"]:
        rows = envelope["sections"][section["id"]]
        if not isinstance(rows, list) or not section["min_rows"] <= len(rows) <= section["max_rows"]: raise AppError("FIELD_INVALID", 422)
        total += len(rows)
        if total > 100: raise AppError("FIELD_INVALID", 422)
        ids = {f["id"] for f in section["row_fields"]}
        for row in rows:
            if not isinstance(row, dict) or set(row) != ids: raise AppError("FIELD_INVALID", 422)
            for f in section["row_fields"]: _field_value(f, row[f["id"]])
    return envelope


def _display(field, value):
    if value is None: return ""
    if field["type"] == "date" and value:
        y,m,d = value.split("-"); return {"YYYY-MM-DD": value, "DD/MM/YYYY": f"{d}/{m}/{y}", "MM/DD/YYYY": f"{m}/{d}/{y}"}[field["display_format"]]
    if field["type"] == "select": return next((x["label"] for x in field["options"] if x["value"] == value), "")
    if field["type"] == "boolean": return field["true_text"] if value is True else field["false_text"] if value is False else ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _alpha(number):
    output = ""
    while number:
        number, rem = divmod(number - 1, 26); output = chr(65 + rem) + output
    return output


def make_context(config, values):
    context = {f["placeholder"]: _display(f, values["fields"][f["id"]]) for f in config["fields"]}
    for section in config["sections"]:
        rows = []
        numbering = section["numbering"]
        for index, raw in enumerate(values["sections"][section["id"]]):
            row = {f["placeholder"]: _display(f, raw[f["id"]]) for f in section["row_fields"]}
            if numbering["style"] != "none":
                number = numbering["start"] + index
                label = str(number) if numbering["style"] == "decimal" else _alpha(number)
                if numbering["style"] == "lower-alpha": label = label.lower()
                row[numbering["placeholder"]] = numbering["prefix"] + label + numbering["suffix"]
            rows.append(row)
        context[section["placeholder"]] = rows
    return context


def safe_filename(pattern, context):
    name = ROOT_TOKEN.sub(lambda m: str(context[m.group(1)]), pattern)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = re.sub(r"_+", "_", name).strip(" .")
    stem = name.rsplit(".", 1)[0]
    if not name or len(name) > 120 or stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1,10)), *(f"LPT{i}" for i in range(1,10))}:
        raise AppError("OUTPUT_NAME_INVALID", 422)
    return name


def generate_docx(docx, config_bytes, values_bytes, expected):
    result = validate_pair(docx, config_bytes)
    if result["fingerprint"] != expected: raise AppError("TEMPLATE_CHANGED", 409)
    values = validate_values(result["configuration"], parse_json(values_bytes, VALUES_MAX, "FIELD_INVALID"))
    context = make_context(result["configuration"], values)
    env = SandboxedEnvironment(undefined=StrictUndefined, autoescape=True)
    env.globals.clear(); env.filters.clear(); env.tests.clear()
    template = DocxTemplate(io.BytesIO(docx))
    try:
        template.render(context, jinja_env=env, autoescape=True)
        for prop in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category"):
            setattr(template.docx.core_properties, prop, "")
        output = io.BytesIO(); template.save(output); rendered = output.getvalue()
        _inspect_zip(rendered)
    except AppError: raise
    except Exception: raise AppError("DOCUMENT_FAILED", 500)
    if len(rendered) > 8_388_608: raise AppError("DOCUMENT_FAILED", 500)
    return rendered, safe_filename(result["configuration"]["output"]["docx_filename_pattern"], context)
def fromstring(data):
    return etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False))


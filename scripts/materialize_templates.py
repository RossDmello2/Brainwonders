"""Materialize the immutable pack's complete schema and fictional examples.

Run from any directory. This is an authoring/build utility, not application code.
"""
import json
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
PACK = ROOT.parent / 'cloud-legal-stenographer-spec-pack'


def strip_default_private_parts(path):
    """python-docx's starter file contains bibliography customXml and a thumbnail."""
    source = path.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as zin, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as zout:
        removed = {'docProps/thumbnail.jpeg', 'customXml/item1.xml', 'customXml/itemProps1.xml', 'customXml/_rels/item1.xml.rels'}
        for info in zin.infolist():
            if info.filename in removed:
                continue
            data = zin.read(info)
            if info.filename.endswith('.rels'):
                root = ET.fromstring(data)
                for child in list(root):
                    if 'customXml' in child.get('Type', '') or 'metadata/thumbnail' in child.get('Type', ''):
                        root.remove(child)
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            elif info.filename == '[Content_Types].xml':
                root = ET.fromstring(data)
                for child in list(root):
                    if child.get('PartName', '').startswith('/customXml/') or child.get('PartName') == '/docProps/thumbnail.jpeg':
                        root.remove(child)
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            zout.writestr(info, data)
    path.write_bytes(output.getvalue())


def fences(name):
    return re.findall(r'```json\n(.*?)\n```', (PACK / name).read_text(encoding='utf-8'), re.S)


def materialize():
    assets = ROOT / 'frontend/templates'
    schema = fences('TEMPLATE_JSON_SCHEMA.md')[0] + '\n'
    for target in [assets / 'schema/template-config.schema.json', ROOT / 'backend/app/schemas/template-config.schema.json']:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(schema, encoding='utf-8', newline='\n')
    guide = fences('TEMPLATE_AUTHORING_GUIDE.md')
    adding = fences('ADDING_OR_REPLACING_TEMPLATES.md')
    assets.mkdir(parents=True, exist_ok=True)
    (assets / 'index.json').write_text(adding[0] + '\n', encoding='utf-8')
    fixture = ROOT / 'tests/fixtures/fictional_values.json'
    fixture.parent.mkdir(parents=True, exist_ok=True)
    fixture.write_text(guide[1] + '\n', encoding='utf-8')
    for raw in [guide[0], adding[1]]:
        config = json.loads(raw)
        folder = assets / config['id'] / config['version']
        folder.mkdir(parents=True, exist_ok=True)
        (folder / 'template.json').write_text(raw + '\n', encoding='utf-8')
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25)
        normal = doc.styles['Normal']
        normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(6)

        def paragraph(block):
            p = doc.add_paragraph()
            p.add_run(block['text'])
            p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                           'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[block['align']]
            return p

        for block in config['preview']['blocks']:
            if block['type'] in ('heading', 'paragraph'):
                paragraph(block)
            elif block['type'] == 'repeat':
                doc.add_paragraph('{%p for row in ' + block['section'] + ' %}')
                for row_block in block['blocks']:
                    paragraph(row_block)
                doc.add_paragraph('{%p endfor %}')
            else:
                columns = block['columns']
                table = doc.add_table(rows=4, cols=len(columns))
                table.autofit = False
                for i, column in enumerate(columns):
                    table.columns[i].width = Mm(160 * column['width_percent'] / 100)
                    for row in table.rows:
                        row.cells[i].width = table.columns[i].width
                    table.cell(0, i).text = column['label']
                    table.cell(2, i).text = column['text']
                table.cell(1, 0).text = '{%tr for row in ' + block['section'] + ' %}'
                table.cell(3, 0).text = '{%tr endfor %}'
        for prop in ['author', 'last_modified_by', 'title', 'subject', 'comments', 'keywords', 'category']:
            setattr(doc.core_properties, prop, '')
        docx_path = folder / 'template.docx'
        doc.save(docx_path)
        strip_default_private_parts(docx_path)


if __name__ == '__main__':
    materialize()
